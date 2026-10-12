#!/usr/bin/env python3
"""Turn the photo archive spreadsheet into the gallery JSON files.

    python scripts/import_sheet.py [path/to/sheet.xlsx]

Reads the "Main" tab (and the "Options" tab for country names) and rewrites the
`photos` list of src/categories/homeland.json and diaspora.json, keeping each
file's other settings. Photos are split by the Location column. Every image URL
is checked; rows that are skipped or look wrong are listed at the end.

Needs openpyxl (pip install openpyxl). The spreadsheet lives in data/, which is
not published and not committed: its Contacts tab holds personal details.
"""

import json
import sys
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SHEET = ROOT / "data" / "Photo_Archive_Final_Output_Website.xlsx"
CATEGORIES = ROOT / "src" / "categories"

# Location column value -> category slug
LOCATIONS = {"Homeland": "homeland", "Diaspora": "diaspora"}
UNKNOWN = {"", "nan"}

# Main tab columns (0-based)
ID, PERIOD, LOCATION, COUNTRY, REGION, LOCALITY, SUBLOCALITY, KEYWORDS, TITLE, DESCRIPTION, S3 = 0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 12


def text(value):
    """Cell value as trimmed text; 'nAn' and empty cells become ''."""
    value = "" if value is None else " ".join(str(value).split())
    return "" if value.lower() in UNKNOWN else value


def country_names(workbook):
    """Country code -> country name, from the Options tab (columns J and K)."""
    names = {}
    for row in workbook["Options"].iter_rows(min_row=2, min_col=10, max_col=11, values_only=True):
        name, code = text(row[0]), text(row[1])
        if name and code:
            names[code] = name
    return names


def url_status(url):
    request = urllib.request.Request(url, method="HEAD")
    try:
        with urllib.request.urlopen(request, timeout=20) as response:
            return response.status
    except urllib.error.HTTPError as err:
        return err.code
    except Exception as err:  # network error, timeout
        return str(err)


def main():
    sheet = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SHEET
    workbook = openpyxl.load_workbook(sheet, data_only=True, read_only=True)
    countries = country_names(workbook)

    rows = [r for r in workbook["Main"].iter_rows(values_only=True) if isinstance(r[ID], (int, float))]
    issues = []
    photos = {slug: [] for slug in LOCATIONS.values()}

    # Raw spaces are not valid in a URL; the CDN expects '+'.
    urls = [text(r[S3]).replace(" ", "+") for r in rows]
    with ThreadPoolExecutor(max_workers=16) as pool:
        statuses = list(pool.map(url_status, urls))

    for row, url, status in zip(rows, urls, statuses):
        photo_id = int(row[ID])
        location = text(row[LOCATION])
        code = text(row[COUNTRY])

        if location not in LOCATIONS:
            issues.append(f"{photo_id}: skipped, no Location (Homeland/Diaspora)")
            continue
        if status != 200:
            issues.append(f"{photo_id}: skipped, image does not load ({status}) {url}")
            continue
        if url != text(row[S3]):
            issues.append(f"{photo_id}: image URL had spaces, fixed to {url}")
        if code and code not in countries:
            issues.append(f"{photo_id}: country code {code!r} not in Options, shown as-is")
        if url.lower().endswith(".bmp"):
            issues.append(f"{photo_id}: image is .bmp (large, should be converted)")
        if not text(row[DESCRIPTION]):
            issues.append(f"{photo_id}: no description")

        tags = []
        for tag in str(row[KEYWORDS] or "").split(","):
            tag = text(tag)
            if tag and tag.lower() not in (t.lower() for t in tags):
                tags.append(tag)

        photos[LOCATIONS[location]].append({
            "id": photo_id,
            "title": text(row[TITLE]),
            "image": url,
            "year": text(row[PERIOD]),
            "country": countries.get(code, code),
            # The most specific place we have.
            "city": text(row[SUBLOCALITY]) or text(row[LOCALITY]) or text(row[REGION]),
            "tags": tags,
            "description": text(row[DESCRIPTION]),
        })

    for slug, items in photos.items():
        file = CATEGORIES / f"{slug}.json"
        title = slug.capitalize()
        category = json.loads(file.read_text(encoding="utf-8")) if file.exists() else {
            "slug": slug,
            "title": title,
            "metaTitle": f"{title} - TXYDE",
            "description": "The Visual Memory of the Circassian People",
            "placesLabel": "Places",
        }
        category["photos"] = sorted(items, key=lambda p: p["id"])
        file.write_text(json.dumps(category, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"{file.relative_to(ROOT)}: {len(items)} photos")

    print(f"\n{len(rows)} rows read, {sum(map(len, photos.values()))} imported.")
    if issues:
        print("\nIssues:")
        print("\n".join(f"  {i}" for i in issues))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
